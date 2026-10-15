import io
import zipfile

import pytest
from docx import Document
from easyresume.documents import MAX_FILE, extract_document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def test_docx_text_and_tables_are_extracted_in_isolation():
    doc = Document()
    doc.add_paragraph("合成简历：本科，Python，3年工作经验")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "合成教育记录"
    output = io.BytesIO()
    doc.save(output)
    result = extract_document(output.getvalue())
    assert result["format"] == "docx"
    assert "合成教育记录" in result["text"] and "Python" in result["text"]


def test_text_pdf_is_extracted():
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 30 100 Td (Synthetic Resume Python) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    assert "Python" in extract_document(output.getvalue())["text"]


def test_scanned_encrypted_and_too_many_pages_are_clear_errors():
    writer = PdfWriter()
    writer.add_blank_page(300, 200)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="扫描件"):
        extract_document(output.getvalue())
    writer.encrypt("synthetic-password")
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="加密"):
        extract_document(output.getvalue())
    writer = PdfWriter()
    for _ in range(51):
        writer.add_blank_page(100, 100)
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="50页"):
        extract_document(output.getvalue())


@pytest.mark.parametrize("size", [0, MAX_FILE + 1], ids=["empty", "over-20mb"])
def test_upload_size_bound(size):
    with pytest.raises(ValueError, match="20MB"):
        extract_document(b"x" * size)


def test_zip_is_not_treated_as_docx_and_macro_rejected():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("word/document.xml", "synthetic")
        archive.writestr("word/vbaProject.bin", "synthetic")
    with pytest.raises(ValueError, match="无宏"):
        extract_document(output.getvalue())


def test_txt_and_doc_conversion():
    assert "本科" in extract_document("合成简历 本科 Python".encode())["text"]
    with pytest.raises(ValueError, match="旧版DOC"):
        extract_document(b"\xd0\xcf\x11\xe0synthetic")
