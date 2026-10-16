"""Anonymous regression examples: no user resume or captured production page is used."""

import io

import pytest
from docx import Document
from easyresume.documents import extract_document
from easyresume.engine import evaluate
from easyresume.parsing import parse_job, parse_profile, skills_from
from easyresume.resume import normalize_profile, summarize_experience
from easyresume.schemas import EducationRecord, Profile, SkillRecord, Strategy, WorkRecord
from pydantic import ValidationError


def checked_profile(**kwargs):
    return Profile(schema_version=2, confirmed=True, experience_as_of="2026-10", **kwargs)


def education(record_id="education", **kwargs):
    values = {
        "id": record_id,
        "degree": "本科",
        "status": "completed",
        "start_date": "2019-09",
        "end_date": "2023-06",
        "confirmed": True,
    }
    values.update(kwargs)
    return EducationRecord(**values)


def work(record_id="work", **kwargs):
    values = {
        "id": record_id,
        "company": "合成公司",
        "start_date": "2020-01",
        "end_date": "2023-01",
        "employment_type": "full_time",
        "confirmed": True,
    }
    values.update(kwargs)
    return WorkRecord(**values)


def checked_job(text):
    job = parse_job(text, title="运维工程师", company="合成企业")
    job.extraction_checked = True
    return job


def test_docx_preserves_interleaved_blocks_and_unique_merged_cells():
    document = Document()
    document.add_paragraph("教育经历")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "2021-10 ~ 2024-07"
    table.cell(1, 0).text = "合成职业学院"
    table.cell(1, 1).text = "计算机网络技术（大专）"
    document.add_paragraph("工作经历")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "合成科技公司"
    table.cell(0, 1).text = "2024-08 ~ 至今"
    table.cell(1, 1).text = "全职运维工程师"
    document.add_paragraph("专业技能\n熟悉Linux系统")
    binary = io.BytesIO()
    document.save(binary)
    extracted = extract_document(binary.getvalue())
    assert extracted["text"].index("合成职业学院") < extracted["text"].index("工作经历")
    assert extracted["text"].count("2021-10 ~ 2024-07") == 1
    assert extracted["text"].count("合成科技公司") == 1
    assert all("location" in block for block in extracted["blocks"])
    parsed = parse_profile(extracted["text"])
    assert parsed.education_records[0].end_date == "2024-07"
    assert parsed.work_records[0].current is True
    assert parsed.work_records[0].company == "合成科技公司"


def test_education_dates_and_cross_line_current_status_stay_with_record():
    value = parse_profile(
        "教育经历\n2019.09-2022.06\n合成学院 大专 全日制\n2023.09-2027.06\n另一大学 本科\n在读 非全日制"
    )
    assert len(value.education_records) == 2
    completed, current = value.education_records
    assert (completed.start_date, completed.end_date, completed.degree, completed.status) == (
        "2019-09",
        "2022-06",
        "大专",
        "completed",
    )
    assert (current.start_date, current.end_date, current.expected_graduation_date, current.status) == (
        "2023-09",
        "",
        "2027-06",
        "current",
    )
    assert not completed.confirmed and not current.confirmed


def test_degree_word_in_work_description_cannot_upgrade_education():
    value = parse_profile(
        "教育经历\n2020-2023\n合成学院 大专\n工作经历\n2023.07-2025.07\n合成公司\n负责本科教学系统维护"
    )
    assert value.education_completed == "大专"
    assert len(value.education_records) == 1
    assert value.education_records[0].start_date == "2020"
    assert value.education_records[0].end_date == "2023"


def test_next_company_before_date_stays_with_its_own_record():
    value = parse_profile(
        "工作经历\n合成甲公司\n2020.01-2022.01\n全职运维工程师\n合成乙公司\n2022.02-2024.02\n全职网络工程师"
    )
    assert [record.company for record in value.work_records] == ["合成甲公司", "合成乙公司"]


def test_skills_handle_chinese_boundaries_negation_and_specific_products():
    text = "熟悉Linux系统，使用Excel表格，不熟悉Python，计划学习Docker，维护MySQL和PostgreSQL，不熟悉Java但熟练TypeScript"
    skills = skills_from(text)
    assert set(skills) == {"Linux", "Excel", "MySQL", "PostgreSQL", "TypeScript"}
    records = {r.name: r for r in parse_profile(text).skill_records}
    assert records["Python"].usage == "negated"
    assert records["Docker"].usage == "learning"
    assert records["MySQL"].category == "数据库"
    assert skills_from("TypeScript JavaScript SQL Server") == ["JavaScript", "TypeScript", "SQL Server"]


def test_profile_summaries_use_only_confirmed_records():
    value = checked_profile(
        education_completed="博士",
        experience_years=50,
        skills=["Invented"],
        education_records=[education(), education("unchecked", degree="博士", confirmed=False)],
        work_records=[work()],
        skill_records=[
            SkillRecord(id="python", name="Python", confirmed=True, usage="used"),
            SkillRecord(id="java", name="Java", usage="used"),
            SkillRecord(id="docker", name="Docker", confirmed=True, usage="learning"),
        ],
    )
    value = normalize_profile(value)
    assert value.education_completed == "本科"
    assert value.experience_years == 3
    assert value.skills == ["Python"]


def test_legacy_summary_is_not_reparsed_or_replaced():
    legacy = Profile(education_completed="本科", experience_years=7, skills=["Custom Skill"], confirmed=True)
    assert normalize_profile(legacy).model_dump() == legacy.model_dump()


def test_manual_skill_edit_is_authoritative_over_original_text():
    value = checked_profile(
        text="熟悉Linux和Python",
        skill_records=[SkillRecord(id="edited", name="PowerShell", source="manual", confirmed=True)],
        skills=["Linux", "Python"],
    )
    assert normalize_profile(value).skills == ["PowerShell"]
    value.skill_records[0].confirmed = False
    assert normalize_profile(value).skills == []


def test_skill_summary_cap_is_explicit_and_full_records_are_preserved():
    value = checked_profile(
        skill_records=[SkillRecord(id=str(i), name=f"Skill {i}", confirmed=True) for i in range(45)]
    )
    normalized = normalize_profile(value)
    assert len(normalized.skills) == 40
    assert len(normalized.skill_records) == 45
    assert any("40" in note for note in normalized.review_notes)


def test_overlapping_formal_work_is_unioned_and_internship_separate():
    value = checked_profile(
        work_records=[
            work("a"),
            work("b", start_date="2021-01", end_date="2024-01"),
            work("intern", start_date="2019-01", end_date="2020-01", employment_type="internship"),
        ]
    )
    result = summarize_experience(value)
    assert result["months"] == 48
    assert result["included_ids"] == ["a", "b"]
    assert result["excluded"][0]["id"] == "intern"
    assert summarize_experience(value, ["intern"])["months"] == 12


@pytest.mark.parametrize(
    "changes",
    [
        {"start_date": "2020"},
        {"employment_type": "unknown"},
        {"confirmed": False},
        {"start_date": "2024-01", "end_date": "2020-01"},
        {"end_date": "2027-01"},
        {"current": True, "end_date": "2023-01"},
    ],
)
def test_incomplete_or_conflicting_work_is_unknown_not_zero(changes):
    result = summarize_experience(checked_profile(work_records=[work(**changes)]))
    assert result["months"] is None
    assert result["excluded"]


def test_current_work_uses_frozen_as_of_and_complete_month_difference():
    value = checked_profile(work_records=[work(start_date="2025-09", end_date="", current=True)])
    assert summarize_experience(value)["months"] == 13
    assert summarize_experience(value, ["missing-id"])["months"] is None
    assert summarize_experience(value, [])["months"] == 0


def test_fulltime_form_cannot_be_borrowed_from_lower_degree():
    value = checked_profile(
        education_records=[
            education("associate", degree="大专", education_form="full_time"),
            education("bachelor", education_form="part_time"),
        ]
    )
    result = evaluate(checked_job("全日制本科"), value, Strategy())
    assert result["group"] == "stretch"
    assert evaluate(checked_job("全日制本科"), value, Strategy(strict_education=True))["group"] == "excluded"
    assert evaluate(checked_job("本科"), value, Strategy())["group"] == "candidate"


def test_campus_year_degree_and_form_must_come_from_same_record():
    value = checked_profile(
        education_records=[
            education(
                "associate",
                degree="大专",
                end_date="2027-06",
                status="current",
                expected_graduation_date="2027-06",
            ),
            education("bachelor", end_date="2023-06"),
        ]
    )
    value.education_records[0].end_date = ""
    assert evaluate(checked_job("2027届本科"), value, Strategy())["group"] == "stretch"
    value.education_records[0].degree = "本科"
    assert evaluate(checked_job("2027届本科"), value, Strategy())["group"] == "candidate"


def test_future_completed_education_requires_review():
    value = checked_profile(education_records=[education(end_date="2028-06")])
    assert normalize_profile(value).education_completed == "unknown"
    assert evaluate(checked_job("本科"), value, Strategy())["group"] == "pending"


def test_only_selected_relevant_experience_qualifies():
    value = checked_profile(
        education_records=[education()],
        work_records=[work("a"), work("b", start_date="2025-01", end_date="2026-01")],
    )
    job = checked_job("本科，2年相关工作经验")
    assert evaluate(job, value, Strategy())["group"] == "pending"
    assert evaluate(job, value, Strategy(relevant_work_ids=["b"]))["group"] == "stretch"
    assert evaluate(job, value, Strategy(relevant_work_ids=["a"]))["group"] == "candidate"
    assert evaluate(job, value, Strategy(relevant_work_ids=["missing"]))["group"] == "pending"


def test_experience_alternative_can_satisfy_degree_requirement():
    value = checked_profile(education_records=[education(degree="大专")], work_records=[work()])
    assert (
        evaluate(checked_job("本科或3年相关经验"), value, Strategy(relevant_work_ids=["work"]))["group"]
        == "candidate"
    )


def test_duplicate_record_ids_and_invalid_months_are_rejected():
    with pytest.raises(ValidationError):
        checked_profile(work_records=[work(), work()])
    with pytest.raises(ValidationError):
        work(start_date="2024-13")
