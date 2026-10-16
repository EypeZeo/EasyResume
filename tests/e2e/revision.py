"""Revision acceptance on synthetic input; AI transport is mocked, no external requests."""

import io
import json
import os
from pathlib import Path

import httpx
from docx import Document
from playwright.sync_api import expect, sync_playwright

base = os.environ.get("EASYRESUME_TEST_URL", "http://127.0.0.1:8773")
headers = {"Origin": base, "X-EasyResume-Client": "workbench"}
headers["Authorization"] = "Bearer " + httpx.post(base + "/api/v1/session", headers=headers).json()["token"]
output = Path("artifacts/screenshots")
output.mkdir(parents=True, exist_ok=True)
document = Document()
document.add_paragraph("姓名：合成样例人")
document.add_paragraph("邮箱：synthetic@example.test")
document.add_paragraph("教育经历")
table = document.add_table(rows=1, cols=2)
table.cell(0, 0).text = "2020.09-2024.06"
table.cell(0, 1).text = "示例学院\n计算机科学 本科 全日制"
document.add_paragraph("工作经历")
document.add_paragraph("2024.07-2025.07 示例科技有限公司 运维工程师 正式工作")
document.add_paragraph("使用Linux系统，维护MySQL数据库，使用Citrix。")
document.add_paragraph("专业技能")
document.add_paragraph("熟悉Linux系统和MySQL，不熟悉Python。")
content = io.BytesIO()
document.save(content)

with sync_playwright() as p:
    options = {} if os.environ.get("EASYRESUME_TEST_BROWSER") == "chromium" else {"channel": "chrome"}
    browser = p.chromium.launch(headless=True, **options)
    page = browser.new_page(viewport={"width": 1440, "height": 1050}, locale="zh-CN")
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto(base)
    page.get_by_role("button", name="我的简历", exact=True).click()
    upload = page.get_by_label("上传简历", exact=True)
    expect(upload).to_be_enabled()
    file = {
        "name": "synthetic-resume.docx",
        "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "buffer": content.getvalue(),
    }
    upload.set_input_files(file)
    suggestions = page.get_by_role("region", name="识别建议")
    expect(suggestions).to_be_visible(timeout=35000)
    page.get_by_role("button", name="采用所选建议", exact=True).click()
    expect(page.get_by_label("入学时间", exact=True).first).to_have_value("2020-09")
    expect(page.get_by_label("结束时间", exact=True).first).to_have_value("2024-06")
    page.get_by_label("学校", exact=True).first.fill("人工更正学院")
    page.get_by_label("已核对这段教育经历", exact=True).first.check()
    page.get_by_label("工作类型", exact=True).first.select_option("full_time")
    page.get_by_label("已核对这段工作经历", exact=True).first.check()
    expect(page.get_by_text("正式工作约 1 年 0 个月", exact=False)).to_be_visible()
    page.get_by_label("简历名称", exact=True).fill("合成经历回归简历")
    expect(page.get_by_text("草稿已保存", exact=True)).to_be_visible(timeout=10000)
    page.get_by_role("button", name="岗位工作台", exact=False).click()
    page.get_by_role("button", name="我的简历", exact=True).click()
    expect(page.get_by_label("学校", exact=True).first).to_have_value("人工更正学院")
    page.reload()
    page.get_by_role("button", name="我的简历", exact=True).click()
    expect(page.get_by_label("学校", exact=True).first).to_have_value("人工更正学院")
    # A repeated parse is a proposal, not an overwrite of the correction.
    page.get_by_role("button", name="重新识别", exact=True).click()
    expect(suggestions).to_be_visible()
    expect(page.get_by_label("采用教育经历建议", exact=False)).not_to_be_checked()
    page.get_by_role("button", name="保留当前内容", exact=True).click()
    expect(page.get_by_label("学校", exact=True).first).to_have_value("人工更正学院")
    page.get_by_role("button", name="保存简历版本", exact=True).click()
    expect(page.get_by_text("简历已保存", exact=True).first).to_be_visible(timeout=10000)
    page.screenshot(path=str(output / "revision-resume.png"), full_page=True)
    saved = httpx.get(base + "/api/v1/state", headers=headers).json()["profiles"][0]
    assert saved["data"]["education_completed"] == "本科"
    assert saved["data"]["graduation_year"] == 2024
    assert saved["data"]["experience_years"] == 1
    # Same file may be selected again, and rejected without losing the current resume.
    page.get_by_label("上传简历", exact=True).set_input_files(file)
    expect(suggestions).to_be_visible(timeout=35000)
    page.get_by_role("button", name="保留当前内容", exact=True).click()
    page.get_by_role("button", name="筛选策略", exact=True).click()
    page.get_by_label("策略名称", exact=True).fill("尚未提交的合成策略")
    expect(page.get_by_text("草稿已保存", exact=True)).to_be_visible(timeout=10000)
    page.get_by_role("button", name="设置与连接", exact=True).click()
    page.get_by_role("button", name="筛选策略", exact=True).click()
    expect(page.get_by_label("策略名称", exact=True)).to_have_value("尚未提交的合成策略")
    page.get_by_role("button", name="设置与连接", exact=True).click()
    page.get_by_role("button", name="数据管理", exact=True).click()
    first = page.get_by_role("button", name="导出备份", exact=True).bounding_box()
    second = page.locator("label.file-button").bounding_box()
    assert first and second and second["x"] - first["x"] - first["width"] >= 11
    backup = httpx.get(base + "/api/v1/backup", headers=headers).json()
    backup_file = {
        "name": "synthetic-backup.json",
        "mimeType": "application/json",
        "buffer": json.dumps(backup).encode(),
    }
    for _ in range(2):
        page.get_by_label("从备份恢复", exact=True).set_input_files(backup_file)
        modal = page.get_by_role("dialog", name="恢复前检查")
        expect(modal.get_by_role("columnheader", name="已存在，跳过")).to_be_visible()
        expect(modal.locator("pre")).to_have_count(0)
        modal.get_by_role("button", name="关闭对话框", exact=True).click()
    page.screenshot(path=str(output / "revision-settings.png"), full_page=True)
    page.get_by_role("button", name="关于与帮助", exact=True).click()
    page.get_by_label("启用 AI 辅助识别", exact=True).check()
    page.get_by_label("接口地址", exact=True).fill("https://example.invalid/v1")
    page.get_by_role("button", name="保存 AI 设置", exact=True).click()
    expect(page.get_by_text("AI 设置已保存", exact=True)).to_be_visible()
    page.get_by_role("button", name="我的简历", exact=True).click()
    # Browser intercept prevents all actual provider calls in this test.
    sent = []

    def mock_task(route):
        if route.request.method == "POST":
            sent.append(route.request.post_data_json)
            route.fulfill(json={"id": "mock-ai", "status": "waiting"})
        elif route.request.method == "DELETE":
            route.fulfill(json={"id": "mock-ai", "status": "cancelled"})
        else:
            route.fulfill(json={"id": "mock-ai", "status": "waiting"})

    page.route("**/api/v1/ai/tasks**", mock_task)
    page.get_by_role("button", name="AI 辅助识别", exact=True).click()
    ai_preview = page.get_by_role("region", name="AI 发送预览")
    expect(ai_preview).to_be_visible()
    preview_text = ai_preview.get_by_role("textbox").input_value()
    assert "合成样例人" not in preview_text and "synthetic@" not in preview_text
    assert not sent
    page.get_by_role("button", name="确认发送脱敏文本", exact=True).click()
    expect(page.get_by_role("button", name="取消等待", exact=True)).to_be_visible()
    assert sent[0]["text"] == preview_text
    page.get_by_role("button", name="取消等待", exact=True).click()
    expect(page.get_by_label("学校", exact=True).first).to_have_value("人工更正学院")
    page.set_viewport_size({"width": 390, "height": 844})
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), "Narrow resume overflows"
    page.screenshot(path=str(output / "revision-mobile.png"), full_page=True)
    page.set_viewport_size({"width": 1440, "height": 1050})
    page.evaluate("document.documentElement.style.zoom = '2'")
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), "200% resume overflows"
    assert not errors, errors
    print(
        json.dumps(
            {
                "result": "PASS",
                "checks": [
                    "DOCX ordered education dates",
                    "work-month union",
                    "manual correction",
                    "draft navigation/reload",
                    "reparse proposal",
                    "save structured profile",
                    "same-file selection",
                    "strategy draft",
                    "backup spacing and counts",
                    "AI redacted preview+confirmation+cancel (mock)",
                    "390px and 200% layout",
                ],
                "real_ai_provider_verified": False,
            },
            ensure_ascii=False,
        )
    )
    browser.close()
